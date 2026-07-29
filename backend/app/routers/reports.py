from typing import Any, List
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import case, func
from io import BytesIO, StringIO
from html import escape
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.core.database import get_db
from app.core.deps import get_current_user, RoleChecker
from app.models.user import User
from app.models.crm import Lead, Deal, AuditLog

router = APIRouter(prefix="/reports", tags=["reports"])

def get_report_data(db: Session, current_user: User):
    # Retrieve data aggregates for reports
    # 1. Lead Conversion by Industry
    industry_query = (
        db.query(
            Lead.industry,
            func.count(Lead.id).label("total_leads"),
            func.sum(case((Lead.status == "Qualified", 1), else_=0)).label("qualified_leads")
        )
        .group_by(Lead.industry)
        .all()
    )
    industry_report = []
    for ind, total, qual in industry_query:
        name = ind or "Unknown"
        rate = (qual / total * 100) if total > 0 else 0
        industry_report.append({
            "industry": name,
            "total_leads": total,
            "qualified_leads": qual,
            "conversion_rate": round(rate, 2)
        })

    # 2. Sales Representative Deal Metrics
    rep_query = (
        db.query(
            User.full_name,
            func.count(Deal.id).label("total_deals"),
            func.sum(Deal.value).label("total_value"),
            func.sum(case((Deal.stage == "Won", Deal.value), else_=0.0)).label("won_value"),
            func.sum(case((Deal.stage == "Won", 1), else_=0)).label("won_count")
        )
        .join(Deal, Deal.owner_id == User.id)
        .group_by(User.full_name)
        .all()
    )
    rep_report = []
    for name, total, val, won_val, won_cnt in rep_query:
        rep_report.append({
            "representative": name,
            "total_deals": total,
            "total_value": val or 0.0,
            "won_value": won_val or 0.0,
            "won_count": won_cnt,
            "win_rate": round((won_cnt / total * 100), 2) if total > 0 else 0.0
        })

    # 3. Marketing Lead Sources Performance
    source_query = (
        db.query(
            Lead.source,
            func.count(Lead.id).label("total_leads"),
            func.count(Deal.id).label("total_deals"),
            func.sum(case((Deal.stage == "Won", Deal.value), else_=0.0)).label("won_revenue")
        )
        .outerjoin(Deal, Deal.lead_id == Lead.id)
        .group_by(Lead.source)
        .all()
    )
    source_report = []
    for src, leads, deals, rev in source_query:
        source_report.append({
            "source": src or "Unknown",
            "total_leads": leads,
            "total_deals": deals,
            "revenue": rev or 0.0
        })

    return {
        "industry_report": industry_report,
        "representative_report": rep_report,
        "source_report": source_report
    }

def log_report_export(db: Session, current_user: User, export_type: str, rows: int) -> None:
    db.add(AuditLog(
        user_id=current_user.id,
        query_type="Report Export",
        raw_query=f"GET /reports/export/{export_type}",
        result_summary=f"Exported report with {rows} aggregate rows"
    ))
    db.commit()

@router.get("/")
def get_reports(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    # Only Admin and Manager can access full operational reports
    RoleChecker(["Admin", "Manager"])(current_user)
    return get_report_data(db, current_user)

@router.get("/export/excel")
def export_reports_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    RoleChecker(["Admin", "Manager"])(current_user)
    data = get_report_data(db, current_user)
    row_count = (
        len(data["representative_report"])
        + len(data["industry_report"])
        + len(data["source_report"])
    )

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    title_font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=11)
    
    header_fill = PatternFill(start_color="8B5CF6", end_color="8B5CF6", fill_type="solid")
    title_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Sheet 1: Sales Performance
    ws1 = wb.create_sheet(title="Sales Performance")
    ws1.append(["Sales Representative Performance Report"])
    ws1.merge_cells("A1:F1")
    ws1.cell(1, 1).font = title_font
    ws1.cell(1, 1).fill = title_fill
    ws1.cell(1, 1).alignment = Alignment(horizontal="center")
    
    ws1.append([]) # empty row
    ws1.append(["Representative", "Total Deals", "Pipeline Value", "Won Value", "Won Count", "Win Rate (%)"])
    for col in range(1, 7):
        cell = ws1.cell(3, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for rep in data["representative_report"]:
        row_idx = ws1.max_row + 1
        ws1.append([
            rep["representative"],
            rep["total_deals"],
            rep["total_value"],
            rep["won_value"],
            rep["won_count"],
            rep["win_rate"]
        ])
        for col in range(1, 7):
            cell = ws1.cell(row_idx, col)
            cell.font = data_font
            cell.border = thin_border
            if col in [3, 4]:
                cell.number_format = "₹#,##0.00"
            elif col in [2, 5, 6]:
                cell.alignment = Alignment(horizontal="right")

    # Sheet 2: Industry Breakdown
    ws2 = wb.create_sheet(title="Industry Breakdown")
    ws2.append(["Lead Conversion by Industry Sector"])
    ws2.merge_cells("A1:D1")
    ws2.cell(1, 1).font = title_font
    ws2.cell(1, 1).fill = title_fill
    ws2.cell(1, 1).alignment = Alignment(horizontal="center")
    
    ws2.append([])
    ws2.append(["Industry", "Total Leads", "Qualified Leads", "Conversion Rate (%)"])
    for col in range(1, 5):
        cell = ws2.cell(3, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for ind in data["industry_report"]:
        row_idx = ws2.max_row + 1
        ws2.append([
            ind["industry"],
            ind["total_leads"],
            ind["qualified_leads"],
            ind["conversion_rate"]
        ])
        for col in range(1, 5):
            cell = ws2.cell(row_idx, col)
            cell.font = data_font
            cell.border = thin_border
            if col in [2, 3, 4]:
                cell.alignment = Alignment(horizontal="right")

    # Sheet 3: Lead Sources
    ws3 = wb.create_sheet(title="Lead Sources")
    ws3.append(["Marketing Lead Sources Effectiveness"])
    ws3.merge_cells("A1:D1")
    ws3.cell(1, 1).font = title_font
    ws3.cell(1, 1).fill = title_fill
    ws3.cell(1, 1).alignment = Alignment(horizontal="center")
    
    ws3.append([])
    ws3.append(["Source Channel", "Total Leads Ingested", "Total Deals Created", "Revenue Generated (₹)"])
    for col in range(1, 5):
        cell = ws3.cell(3, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for src in data["source_report"]:
        row_idx = ws3.max_row + 1
        ws3.append([
            src["source"],
            src["total_leads"],
            src["total_deals"],
            src["revenue"]
        ])
        for col in range(1, 5):
            cell = ws3.cell(row_idx, col)
            cell.font = data_font
            cell.border = thin_border
            if col == 4:
                cell.number_format = "₹#,##0.00"
            elif col in [2, 3]:
                cell.alignment = Alignment(horizontal="right")

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to buffer
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    log_report_export(db, current_user, "excel", row_count)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=crm_operational_report.xlsx"}
    )

@router.get("/export/html")
def export_reports_html(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
) -> Any:
    RoleChecker(["Admin", "Manager"])(current_user)
    data = get_report_data(db, current_user)
    row_count = (
        len(data["representative_report"])
        + len(data["industry_report"])
        + len(data["source_report"])
    )
    representative_rows = "".join(
        "<tr>"
        f"<td>{escape(str(r['representative']))}</td>"
        f"<td class='text-right'>{r['total_deals']}</td>"
        f"<td class='text-right'>₹{r['total_value']:,.2f}</td>"
        f"<td class='text-right'>₹{r['won_value']:,.2f}</td>"
        f"<td class='text-right'>{r['won_count']}</td>"
        f"<td class='text-right'>{r['win_rate']}%</td>"
        "</tr>"
        for r in data["representative_report"]
    )
    industry_rows = "".join(
        "<tr>"
        f"<td>{escape(str(i['industry']))}</td>"
        f"<td class='text-right'>{i['total_leads']}</td>"
        f"<td class='text-right'>{i['qualified_leads']}</td>"
        f"<td class='text-right'>{i['conversion_rate']}%</td>"
        "</tr>"
        for i in data["industry_report"]
    )
    source_rows = "".join(
        "<tr>"
        f"<td>{escape(str(s['source']))}</td>"
        f"<td class='text-right'>{s['total_leads']}</td>"
        f"<td class='text-right'>{s['total_deals']}</td>"
        f"<td class='text-right'>₹{s['revenue']:,.2f}</td>"
        "</tr>"
        for s in data["source_report"]
    )
    
    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <title>CRM Operational Performance Report</title>
        <style>
            body {{ font-family: 'Segoe UI', sans-serif; margin: 40px; color: #1f2937; }}
            h1 {{ border-bottom: 2px solid #8b5cf6; padding-bottom: 10px; color: #111827; }}
            h2 {{ color: #4b5563; margin-top: 40px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; }}
            th, td {{ border: 1px solid #e5e7eb; padding: 12px; text-align: left; }}
            th {{ background-color: #f9fafb; font-weight: bold; }}
            .text-right {{ text-align: right; }}
            .footer {{ margin-top: 60px; font-size: 12px; color: #9ca3af; text-align: center; }}
        </style>
    </head>
    <body>
        <h1>CRM Performance Summary Report</h1>
        <p>Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        
        <h2>1. Sales Representative Performance</h2>
        <table>
            <tr>
                <th>Representative</th>
                <th class="text-right">Total Deals</th>
                <th class="text-right">Pipeline Value</th>
                <th class="text-right">Won Value</th>
                <th class="text-right">Won Count</th>
                <th class="text-right">Win Rate</th>
            </tr>
            {representative_rows}
        </table>

        <h2>2. Industry Conversion Breakdown</h2>
        <table>
            <tr>
                <th>Industry</th>
                <th class="text-right">Total Leads</th>
                <th class="text-right">Qualified Leads</th>
                <th class="text-right">Conversion Rate</th>
            </tr>
            {industry_rows}
        </table>

        <h2>3. Marketing Lead Channels Performance</h2>
        <table>
            <tr>
                <th>Source Channel</th>
                <th class="text-right">Total Leads</th>
                <th class="text-right">Total Deals</th>
                <th class="text-right">Revenue Generated</th>
            </tr>
            {source_rows}
        </table>

        <div class="footer">
            <p>Enterprise CRM Dashboard Analytical Export. Print to save as PDF.</p>
        </div>
    </body>
    </html>
    """
    log_report_export(db, current_user, "html", row_count)
    return StreamingResponse(
        StringIO(html_content),
        media_type="text/html",
        headers={"Content-Disposition": "attachment; filename=crm_performance_report.html"}
    )
